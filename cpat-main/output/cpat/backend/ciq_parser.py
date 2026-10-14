"""
Unified CIQ XLSX parser for RHOSO and OCP Customer Information Questionnaire files.

Auto-detects cluster type (RHOSO/OCP) from Excel table names and Type parameter.
Handles parameter tables (2-column name/value) and row tables (multi-column records).

Supports both RHOSO and OCP in a single workbook (combined CIQ) or separate
workbooks. Each cluster type is returned as a CiqParseResult in a CiqParseBatchResult.
"""

from __future__ import annotations

import json
import os
import re
from dataclasses import dataclass, field
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import range_boundaries


# ---------------------------------------------------------------------------
# Table registries — detail (registration), supplemental, row
# ---------------------------------------------------------------------------

# Detail tables contain the Type parameter → produce a cluster registration
DETAIL_TABLES: dict[str, tuple[str, str]] = {
    "rhoso_details": ("rhoso", "RHOSO"),
    "ocp_details": ("ocp", "OCP"),
    "ocp_cluster_details": ("ocp", "OCP"),
}

# Supplemental parameter tables per cluster type (no Type param, merged into details)
SUPPLEMENTAL_TABLES: dict[str, tuple[str, ...]] = {
    "rhoso": (
        "rhoso_access",
        "rhoso_disconnected",
        "rhoso_storage",
        "rhoso_ceph_storage",
        "rhoso_control_plane",
        "rhoso_services",
        "rhoso_observability",
        "rhoso_pv_sizing",
        "rhoso_cinder_backend",
        "rhoso_neutron_config",
        "rhoso_edpm_network_config",
        "rhoso_data_plane",
        "rhoso_trident_storage",
        "rhoso_lso_storage",
        "rhoso_multipath",
    ),
    "ocp": (
        "ocp_access",
        "ocp_networking",
        "ocp_platform_baremetal",
        "ocp_agent_config",
        "ocp_disconnected",
        "ocp_mirror_registry",
        "ocp_bastion",
        "ocp_load_balancer",
        "ocp_certificates",
        "ocp_storage",
        "ocp_day2_operators",
        "ocp_ingress",
        "ocp_proxy",
    ),
}

ALL_SUPPLEMENTAL_TABLES: frozenset[str] = frozenset(
    table for tables in SUPPLEMENTAL_TABLES.values() for table in tables
)

ALL_PARAM_TABLES: frozenset[str] = frozenset(
    list(DETAIL_TABLES.keys()) + list(ALL_SUPPLEMENTAL_TABLES)
)

# Row tables: table_name → (cluster_type, row_key)
ROW_TABLES: dict[str, tuple[str, str]] = {
    "rhoso_data_plane_nodes": ("rhoso", "nodes"),
    "rhoso_networks": ("rhoso", "networks"),
    "ocp_bmc_hosts": ("ocp", "bmc_hosts"),
    "ocp_nmstate_interfaces": ("ocp", "nmstate_interfaces"),
    "ocp_compute_nodes": ("ocp", "compute_nodes"),
    "ocp_dns_records": ("ocp", "dns_records"),
}

CLUSTER_TYPE_LABELS: dict[str, str] = {
    "rhoso": "RHOSO",
    "ocp": "OCP",
}

# Type parameter value patterns → (cluster_type, label)
TYPE_VALUE_PATTERNS: list[tuple[re.Pattern[str], str, str]] = [
    (
        re.compile(
            r"^(rhoso|red\s*hat\s*openstack\s*services\s*on\s*openshift)$",
            re.IGNORECASE,
        ),
        "rhoso",
        "RHOSO",
    ),
    (
        re.compile(
            r"^(ocp|openshift\s*container\s*platform)$",
            re.IGNORECASE,
        ),
        "ocp",
        "OCP",
    ),
]


# ---------------------------------------------------------------------------
# Header / cell patterns
# ---------------------------------------------------------------------------

PARAM_NAME_HEADER_REGEX = re.compile(r"parameter\s*name", re.IGNORECASE)
PARAM_VALUE_HEADER_REGEX = re.compile(r"parameter\s*value", re.IGNORECASE)
TYPE_PARAM_REGEX = re.compile(r"^type$", re.IGNORECASE)
CLUSTER_NAME_PARAM_REGEX = re.compile(r"cluster\s*name", re.IGNORECASE)

_VERSION_PARAM_REGEX = re.compile(
    r"^(rhoso\s*version|rhocp\s*version|nad\s*cni\s*version"
    r"|ocp\s*version|openshift\s*version)$",
    re.IGNORECASE,
)

_CANONICAL_HEADER_WORDS: dict[str, str] = {
    "IP": "IP", "BMC": "BMC", "SSH": "SSH", "VIP": "VIP", "CIDR": "CIDR",
    "MTU": "MTU", "NAD": "NAD", "NIC": "NIC", "VLAN": "VLAN", "MAC": "MAC",
    "NNCP": "NNCP", "DNS": "DNS", "FQDN": "FQDN", "DHCP": "DHCP",
    "DPDK": "DPDK", "OVN": "OVN", "OVS": "OVS", "BMH": "BMH",
    "EE": "EE", "NTP": "NTP", "ID": "ID", "METALLB": "MetalLB",
    "CPU": "CPU", "NUMA": "NUMA", "NFS": "NFS", "TLS": "TLS", "PV": "PV",
    "CR": "CR", "API": "API", "LSO": "LSO", "CSI": "CSI", "SRIOV": "SR-IOV",
    "IPV4": "IPv4", "IPV6": "IPv6", "ISO": "ISO", "CA": "CA",
}


# ---------------------------------------------------------------------------
# Data classes
# ---------------------------------------------------------------------------

@dataclass
class CiqParseResult:
    """Parsed data from a single cluster registration (RHOSO or OCP)."""
    cluster_name: str
    cluster_type: str
    cluster_type_label: str
    source_table: str = ""
    source_sheet: str = ""
    parameters: dict[str, str] = field(default_factory=dict)
    type_raw_value: str = ""
    row_tables: dict[str, list[dict[str, str]]] = field(default_factory=dict)

    # RHOSO backward-compatible properties
    @property
    def nodes(self) -> list[dict[str, str]]:
        return self.row_tables.get("nodes", [])

    @property
    def networks(self) -> list[dict[str, str]]:
        return self.row_tables.get("networks", [])

    # OCP backward-compatible properties
    @property
    def bmc_hosts(self) -> list[dict[str, str]]:
        return self.row_tables.get("bmc_hosts", [])

    @property
    def nmstate_interfaces(self) -> list[dict[str, str]]:
        return self.row_tables.get("nmstate_interfaces", [])

    @property
    def compute_nodes(self) -> list[dict[str, str]]:
        return self.row_tables.get("compute_nodes", [])

    @property
    def dns_records(self) -> list[dict[str, str]]:
        return self.row_tables.get("dns_records", [])

    def parameters_json(self) -> str:
        data: dict[str, Any] = {
            "parameters": self.parameters,
            "type_raw_value": self.type_raw_value,
            "source_table": self.source_table,
            "source_sheet": self.source_sheet,
        }
        data.update(self.row_tables)
        return json.dumps(data)

    def ansible_extra_vars(self) -> dict[str, Any]:
        """Convert parsed CIQ to the flat ocp_*/rhoso_* Ansible variable dict."""
        prefix = "ocp" if self.cluster_type == "ocp" else "rhoso"
        evars: dict[str, Any] = {}
        for param_name, param_value in self.parameters.items():
            key = _param_to_ansible_var(param_name, prefix)
            evars[key] = param_value
        for row_key, rows in self.row_tables.items():
            if rows:
                evars[f"{prefix}_{row_key}"] = rows
        return evars


@dataclass
class CiqParseBatchResult:
    """All clusters parsed from one CIQ workbook."""
    clusters: list[CiqParseResult] = field(default_factory=list)
    table_parameters: dict[str, dict[str, str]] = field(default_factory=dict)

    def get_by_type(self, cluster_type: str) -> CiqParseResult | None:
        for cluster in self.clusters:
            if cluster.cluster_type == cluster_type:
                return cluster
        return None


class CiqParseError(Exception):
    pass


# Backward-compatible alias
OcpCiqParseError = CiqParseError


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _normalize_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value).strip()
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return format(value, "f").rstrip("0").rstrip(".")
    if isinstance(value, int):
        return str(value)
    return str(value).strip()


def _normalize_version_cell(value: Any, cluster_type: str = "") -> str:
    """Normalize version cells.

    RHOSO: Excel numeric 18 / 18.0 → '18.0'
    OCP:   Excel numeric 4.19 → '4.19', bare int 20 → '4.20'
    """
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value).strip()
    if isinstance(value, float):
        if cluster_type == "rhoso" and value.is_integer():
            return f"{int(value)}.0"
        return format(value, "f").rstrip("0").rstrip(".")
    if isinstance(value, int):
        if cluster_type == "ocp":
            return f"4.{value}" if value < 100 else str(value)
        return f"{value}.0"
    text = str(value).strip()
    if re.fullmatch(r"[0-9]+", text):
        if cluster_type == "ocp":
            return f"4.{text}"
        return f"{text}.0"
    return text


_LOWERCASE_CONNECTORS: set[str] = {
    "OR", "AND", "OF", "FOR", "IN", "TO", "BY", "ON", "AT", "THE", "A", "AN",
}


def _normalize_header(raw: str) -> str:
    """Normalize column header to canonical mixed case.

    UPPERCASE Excel headers (NODE NAME, BMC ADDRESS) are converted to
    title-case with known acronyms preserved (Node Name, BMC Address).
    Connectors like 'or', 'and', 'of' stay lowercase (IP or Target).
    """
    words = _normalize_cell(raw).split()
    result: list[str] = []
    for idx, w in enumerate(words):
        upper = w.upper()
        if upper in _CANONICAL_HEADER_WORDS:
            result.append(_CANONICAL_HEADER_WORDS[upper])
            continue
        if idx > 0 and upper in _LOWERCASE_CONNECTORS:
            result.append(w.lower())
            continue
        matched = False
        for acronym, canonical in _CANONICAL_HEADER_WORDS.items():
            if (upper.startswith(acronym) and len(upper) > len(acronym)
                    and upper[len(acronym):].isdigit()):
                result.append(canonical + upper[len(acronym):])
                matched = True
                break
        if not matched:
            result.append(w.capitalize())
    return " ".join(result)


def _param_to_ansible_var(param_name: str, prefix: str = "ocp") -> str:
    """Convert CIQ parameter name to prefixed Ansible variable name.

    E.g. 'Cluster Name' with prefix 'ocp' → 'ocp_cluster_name'
    """
    key = param_name.strip().lower()
    key = re.sub(r"[^a-z0-9]+", "_", key).strip("_")
    if not key.startswith(f"{prefix}_"):
        key = f"{prefix}_{key}"
    return key


def _read_range(ws, ref: str) -> list[list[str]]:
    min_col, min_row, max_col, max_row = range_boundaries(ref)
    rows: list[list[str]] = []
    for row in ws.iter_rows(
        min_row=min_row, max_row=max_row,
        min_col=min_col, max_col=max_col,
        values_only=True,
    ):
        rows.append([_normalize_cell(cell) for cell in row])
    return rows


def _classify_table(table) -> tuple[str, str] | None:
    """Return (canonical_name, kind) or None.

    kind is one of: 'detail', 'supplemental', 'row'
    """
    names = {table.name.lower()}
    display_name = getattr(table, "displayName", None)
    if display_name:
        names.add(str(display_name).lower())

    for canonical in DETAIL_TABLES:
        if canonical in names:
            return (canonical, "detail")
    for canonical in ALL_SUPPLEMENTAL_TABLES:
        if canonical in names:
            return (canonical, "supplemental")
    for canonical in ROW_TABLES:
        if canonical in names:
            return (canonical, "row")
    return None


def _parse_row_table(rows: list[list[str]]) -> list[dict[str, str]]:
    if len(rows) < 2:
        return []
    headers = [_normalize_header(cell) for cell in rows[0]]
    if not any(headers):
        raise CiqParseError("Row table must contain a header row.")
    records: list[dict[str, str]] = []
    for row in rows[1:]:
        if not any(_normalize_cell(cell) for cell in row):
            continue
        record: dict[str, str] = {}
        for col_idx, header in enumerate(headers):
            if not header:
                continue
            value = row[col_idx] if col_idx < len(row) else ""
            record[header] = _normalize_cell(value)
        if record:
            records.append(record)
    return records


def _parse_parameter_rows(rows: list[list[str]]) -> dict[str, str]:
    if not rows:
        return {}
    header = rows[0]
    name_col = None
    value_col = None
    for col_idx, cell in enumerate(header):
        if PARAM_NAME_HEADER_REGEX.search(cell):
            name_col = col_idx
        if PARAM_VALUE_HEADER_REGEX.search(cell):
            value_col = col_idx
    if name_col is None or value_col is None:
        raise CiqParseError(
            "Each parameter table must contain 'Parameter Name' "
            "and 'Parameter Value' columns."
        )
    parameters: dict[str, str] = {}
    for row in rows[1:]:
        if len(row) <= max(name_col, value_col):
            continue
        param_name = row[name_col]
        param_value = row[value_col]
        if not param_name:
            continue
        if _VERSION_PARAM_REGEX.search(param_name):
            param_value = _normalize_version_cell(param_value)
        parameters[param_name] = param_value
    return parameters


def _resolve_cluster_type(type_value: str) -> tuple[str, str]:
    normalized = _normalize_cell(type_value)
    if not normalized:
        raise CiqParseError("Type parameter value is empty.")
    for pattern, cluster_type, label in TYPE_VALUE_PATTERNS:
        if pattern.search(normalized):
            return cluster_type, label
    allowed = ", ".join(label for _, _, label in TYPE_VALUE_PATTERNS)
    raise CiqParseError(
        f"Invalid Type value '{normalized}'. Expected one of: {allowed}."
    )


def _find_parameter_value(
    parameters: dict[str, str], pattern: re.Pattern[str]
) -> str | None:
    for name, value in parameters.items():
        if pattern.search(name):
            return value
    return None


def _get_supplementals_for_type(cluster_type: str) -> tuple[str, ...]:
    return SUPPLEMENTAL_TABLES.get(cluster_type, ())


def _merge_supplemental_parameters(
    base_parameters: dict[str, str],
    all_table_parameters: dict[str, dict[str, str]],
    cluster_type: str,
) -> dict[str, str]:
    """Merge supplemental parameter tables into the base details dict."""
    merged = dict(base_parameters)
    for table_name in _get_supplementals_for_type(cluster_type):
        for key, extra in all_table_parameters.items():
            if key.endswith(f"/{table_name}"):
                for param_name, param_value in extra.items():
                    if param_name in merged and merged[param_name] != param_value:
                        source_sheet = key.split("/", 1)[0]
                        raise CiqParseError(
                            f"Duplicate parameter '{param_name}' on sheet "
                            f"'{source_sheet}' in table '{table_name}' "
                            f"conflicts with an earlier {cluster_type.upper()} table."
                        )
                    merged[param_name] = param_value
    return merged


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------

def parse_ciq_xlsx(file_path: str) -> CiqParseBatchResult:
    """Parse a CIQ Excel workbook and return all cluster registrations.

    Auto-detects RHOSO and OCP from Excel table names and Type parameter values.
    A single workbook can contain both RHOSO and OCP registrations.

    Raises CiqParseError on any validation failure.
    """
    if not os.path.isfile(file_path):
        raise CiqParseError(f"CIQ file not found: {file_path}")
    if not file_path.lower().endswith(".xlsx"):
        raise CiqParseError("CIQ file must be an .xlsx workbook.")

    try:
        workbook = load_workbook(file_path, data_only=True)
    except Exception as exc:
        raise CiqParseError(f"Failed to read XLSX file: {exc}") from exc

    all_table_parameters: dict[str, dict[str, str]] = {}
    # row_data[cluster_type][row_key] = list of records
    row_data: dict[str, dict[str, list[dict[str, str]]]] = {}
    # registrations: (sheet, canonical_name, cluster_type, params, type_raw, label)
    registrations: list[tuple[str, str, str, dict[str, str], str, str]] = []

    for sheet_name in workbook.sheetnames:
        ws = workbook[sheet_name]
        for table in ws.tables.values():
            classification = _classify_table(table)
            if classification is None:
                continue

            canonical_name, kind = classification

            if kind == "row":
                row_cluster_type, row_key = ROW_TABLES[canonical_name]
                try:
                    records = _parse_row_table(_read_range(ws, table.ref))
                except CiqParseError:
                    raise
                except Exception as exc:
                    raise CiqParseError(
                        f"Failed to parse table '{canonical_name}' on sheet "
                        f"'{sheet_name}': {exc}"
                    ) from exc
                row_data.setdefault(row_cluster_type, {})
                row_data[row_cluster_type].setdefault(row_key, [])
                row_data[row_cluster_type][row_key].extend(records)
                continue

            # Parameter table (detail or supplemental)
            rows = _read_range(ws, table.ref)
            table_key = f"{sheet_name}/{canonical_name}"
            try:
                params = _parse_parameter_rows(rows)
            except CiqParseError:
                raise
            except Exception as exc:
                raise CiqParseError(
                    f"Failed to parse table '{canonical_name}' on sheet "
                    f"'{sheet_name}': {exc}"
                ) from exc

            if params:
                all_table_parameters[table_key] = params

            if kind == "supplemental":
                type_in_supplement = _find_parameter_value(params, TYPE_PARAM_REGEX)
                if type_in_supplement:
                    raise CiqParseError(
                        f"Table '{canonical_name}' on sheet '{sheet_name}' must not "
                        f"contain a Type parameter. Put Type only in the details table."
                    )
                continue

            # Detail table — look for Type parameter
            type_raw = _find_parameter_value(params, TYPE_PARAM_REGEX)
            if not type_raw:
                continue

            expected_type, expected_label = DETAIL_TABLES[canonical_name]
            resolved_type, resolved_label = _resolve_cluster_type(type_raw)

            if resolved_type != expected_type:
                raise CiqParseError(
                    f"Table '{canonical_name}' on sheet '{sheet_name}' has Type "
                    f"'{type_raw}' but table '{canonical_name}' requires "
                    f"'{expected_label}'."
                )

            registrations.append((
                sheet_name, canonical_name, resolved_type,
                params, type_raw, resolved_label,
            ))

    workbook.close()

    if not all_table_parameters and not row_data:
        all_known = sorted(set(DETAIL_TABLES.keys()) | ALL_SUPPLEMENTAL_TABLES | set(ROW_TABLES.keys()))
        raise CiqParseError(
            f"No CIQ tables found. Expected one of: {', '.join(all_known)}."
        )

    if not registrations:
        raise CiqParseError(
            "Type parameter not found in any details table. "
            "Add Type = RHOSO or Type = OCP to the appropriate details table."
        )

    seen_types: set[str] = set()
    seen_names: set[str] = set()
    clusters: list[CiqParseResult] = []

    for sheet_name, canonical_name, cluster_type, parameters, type_raw, label in registrations:
        if cluster_type in seen_types:
            raise CiqParseError(
                f"Duplicate registration for cluster type '{label}' "
                f"in CIQ file. Each type may only be defined once."
            )
        seen_types.add(cluster_type)

        cluster_name = _find_parameter_value(parameters, CLUSTER_NAME_PARAM_REGEX)
        if not cluster_name:
            raise CiqParseError(
                f"Cluster Name parameter not found in table '{canonical_name}' "
                f"on sheet '{sheet_name}'."
            )

        if cluster_name in seen_names:
            raise CiqParseError(
                f"Duplicate Cluster Name '{cluster_name}' found across CIQ tables."
            )
        seen_names.add(cluster_name)

        merged_parameters = _merge_supplemental_parameters(
            parameters, all_table_parameters, cluster_type,
        )

        cluster_row_data = row_data.get(cluster_type, {})

        clusters.append(CiqParseResult(
            cluster_name=cluster_name,
            cluster_type=cluster_type,
            cluster_type_label=label,
            source_table=canonical_name,
            source_sheet=sheet_name,
            parameters=merged_parameters,
            type_raw_value=type_raw,
            row_tables=cluster_row_data,
        ))

    clusters.sort(key=lambda c: c.cluster_name.lower())

    return CiqParseBatchResult(
        clusters=clusters,
        table_parameters=all_table_parameters,
    )


def parse_ocp_ciq_xlsx(file_path: str) -> CiqParseResult:
    """Parse an OCP CIQ workbook. Convenience wrapper around parse_ciq_xlsx.

    Returns the single OCP CiqParseResult. Raises CiqParseError if no OCP
    registration found.
    """
    batch = parse_ciq_xlsx(file_path)
    ocp = batch.get_by_type("ocp")
    if ocp is None:
        raise CiqParseError(
            "No OCP registration found. Add Type = OCP to the ocp_details table."
        )
    return ocp
