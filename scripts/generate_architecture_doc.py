#!/usr/bin/env python3
"""Generate SprintWeave Architecture & Deployment Guide (Word .docx)."""

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from pathlib import Path
import datetime

OUT = Path(__file__).resolve().parent.parent.parent / "SprintWeave_Architecture_Guide.docx"

doc = Document()

style = doc.styles["Normal"]
font = style.font
font.name = "Calibri"
font.size = Pt(11)

# ── Title ──────────────────────────────────────────────────────
doc.add_heading("SprintWeave Architecture & Deployment Guide", level=0)
doc.add_paragraph(
    f"Generated: {datetime.datetime.now().strftime('%B %d, %Y')}\n"
    "Confidential — Red Hat Global Services"
)
doc.add_page_break()

# ── TOC placeholder ────────────────────────────────────────────
doc.add_heading("Table of Contents", level=1)
doc.add_paragraph(
    "1. Executive Summary\n"
    "2. Cursor IDE Architecture\n"
    "3. RHDP / OpenShift Architecture\n"
    "4. Side-by-Side Comparison\n"
    "5. Docker Image Strategy (Two-Layer Build)\n"
    "6. End-to-End Deployment Steps\n"
    "7. Docling PDF Processing Pipeline\n"
    "8. LLM Integration (MaaS / LiteLLM)\n"
    "9. Agent Rule & Context Template System\n"
    "10. Cost Analysis\n"
    "11. File Tree — Cursor vs SprintWeave Output"
)
doc.add_page_break()

# ═══════════════════════════════════════════════════════════════
# 1. Executive Summary
# ═══════════════════════════════════════════════════════════════
doc.add_heading("1. Executive Summary", level=1)
doc.add_paragraph(
    "SprintWeave is an AI-driven framework that takes PDF/DOCX input documents "
    "and produces delivery-ready artifacts — source code, Excel workbooks, "
    "deployment manifests, and documentation. It uses Large Language Models "
    "(Claude Opus via MaaS/LiteLLM) to generate the CPAT (Cloud Platform "
    "Automation Tool) codebase for automated RHOSO deployments."
)
doc.add_paragraph(
    "Two execution environments exist:"
)
t = doc.add_table(rows=3, cols=3)
t.style = "Light Grid Accent 1"
t.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, h in enumerate(["", "Cursor IDE", "RHDP / OpenShift"]):
    t.rows[0].cells[i].text = h
data = [
    ["Where it runs", "Developer laptop (Cursor IDE)", "OpenShift cluster (RHDP lab)"],
    ["LLM billing", "$20/month flat (Cursor subscription)", "$4-6/run (MaaS per-token)"],
]
for r, row in enumerate(data, 1):
    for c, val in enumerate(row):
        t.rows[r].cells[c].text = val

# ═══════════════════════════════════════════════════════════════
# 2. Cursor IDE Architecture
# ═══════════════════════════════════════════════════════════════
doc.add_heading("2. Cursor IDE Architecture", level=1)

doc.add_heading("2.1 Overview", level=2)
doc.add_paragraph(
    "In the Cursor IDE, the pipeline is driven by agents/cursor/run_pipeline.py, "
    "which uses the Cursor SDK (Agent.create() + agent.send()) to run each "
    "sub-stage. The model enters a tool loop — reading files, writing files, "
    "and deciding when to stop — all controlled by the LLM."
)

doc.add_heading("2.2 Pipeline Stages", level=2)
doc.add_paragraph(
    "The RHOSO Delivery pipeline has 7 stages:"
)
t = doc.add_table(rows=8, cols=4)
t.style = "Light Grid Accent 1"
for i, h in enumerate(["Stage", "Agent", "Sub-stage", "Output"]):
    t.rows[0].cells[i].text = h
stages = [
    ["RDA-1", "RHOSO DA", "Schema", "rhoso_ciq_schema.py"],
    ["RDA-2", "RHOSO DA", "Generator", "generate_rhoso_ciq_workbooks.py"],
    ["RDA-3", "RHOSO DA", "Workbooks", "*.xlsx workbooks"],
    ["RDE-1", "RHOSO DE", "Ansible", "25 tasks (9 subdirs), 19 templates, defaults"],
    ["RDE-2", "RHOSO DE", "Backend", "app.py, ansible_runner.py, 9 modules"],
    ["RDE-3", "RHOSO DE", "Frontend", "11 React components, api.js, config"],
    ["RDE-4", "RHOSO DE", "Docker", "Dockerfile, entrypoint.sh, bringup.sh, USAGE.md"],
]
for r, row in enumerate(stages, 1):
    for c, val in enumerate(row):
        t.rows[r].cells[c].text = val

doc.add_heading("2.3 Tool Loop Flow", level=2)
doc.add_paragraph(
    "For each stage, Cursor's agent performs the following steps:"
)
doc.add_paragraph(
    "Step 1: read_file .cursor/rules/agent-rhoso-deployment-engineer.md "
    "(full 22,627 chars — the spec listing every file, component, template)", style="List Number"
)
doc.add_paragraph(
    "Step 2: read_file context/templates/rhoso-deployment-engineer.md "
    "(full 45,747 chars — CIQ-to-CR mapping, workflow phases, manifest catalog)", style="List Number"
)
doc.add_paragraph(
    "Step 3: read_file parsed/*.md (Docling output from input PDFs)", style="List Number"
)
doc.add_paragraph(
    "Steps 4-25: write_file — one file per call. The model writes each file "
    "with the full path including subdirectory (e.g. tasks/secrets/generate_osp_secret.yml). "
    "It decides when to stop — no Python-side file count limits.", style="List Number"
)
doc.add_paragraph(
    "Step 26: verify_outputs() — checks 2-3 key files exist. Only logs warnings, "
    "never blocks or retries.", style="List Number"
)

doc.add_heading("2.4 Why Cursor Gets the Right Output", level=2)
doc.add_paragraph(
    "1. The agent reads the full spec (agent rule + context template) via tool "
    "calls — no truncation.\n"
    "2. It writes one file at a time, so each file path is a deliberate decision "
    "guided by the spec tables.\n"
    "3. The model decides when it's done (stops calling write_file), not a "
    "Python-side counter.\n"
    "4. verify_paths only logs — it never forces the model to generate more."
)

doc.add_heading("2.5 Key Files", level=2)
t = doc.add_table(rows=5, cols=3)
t.style = "Light Grid Accent 1"
for i, h in enumerate(["File", "Size", "Purpose"]):
    t.rows[0].cells[i].text = h
files = [
    ["agents/cursor/run_pipeline.py", "~350 lines", "Pipeline runner — sends prompts, streams output, verifies"],
    [".cursor/rules/agent-rhoso-deployment-engineer.md", "22,627 chars (306 lines)", "Agent rule — lists every file, component, template by name"],
    ["context/templates/rhoso-deployment-engineer.md", "45,747 chars (1014 lines)", "Context template — CIQ mapping, workflow phases, code patterns"],
    ["context/pipelines.yaml", "~90 lines", "Pipeline stage definitions with sub_stage and generates fields"],
]
for r, row in enumerate(files, 1):
    for c, val in enumerate(row):
        t.rows[r].cells[c].text = val

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════
# 3. RHDP / OpenShift Architecture
# ═══════════════════════════════════════════════════════════════
doc.add_heading("3. RHDP / OpenShift Architecture", level=1)

doc.add_heading("3.1 Overview", level=2)
doc.add_paragraph(
    "On RHDP (Red Hat Demo Platform), SprintWeave runs as two processes in a "
    "single container on OpenShift:\n\n"
    "  • Streamlit UI (port 8080) — file upload, pipeline selection, job monitoring\n"
    "  • FastAPI backend (port 8000) — job execution, LLM calls, Docling processing\n\n"
    "Persistent storage is a PVC mounted at /data containing job workspaces."
)

doc.add_heading("3.2 Request Flow", level=2)
doc.add_paragraph(
    "1. User uploads PDF(s) via Streamlit UI (port 8080)", style="List Number"
)
doc.add_paragraph(
    "2. UI calls FastAPI POST /v1/jobs (port 8000) with file references + pipeline choice", style="List Number"
)
doc.add_paragraph(
    "3. Worker creates workspace under /data/jobs/<job-id>/workspace/", style="List Number"
)
doc.add_paragraph(
    "4. Docling converts PDF → parsed/*.md (shared converter, loaded once)", style="List Number"
)
doc.add_paragraph(
    "5. For each pipeline stage, worker calls LLM via MaaS/LiteLLM endpoint", style="List Number"
)
doc.add_paragraph(
    "6. RDE stages use one-shot mode: full spec pre-loaded in prompt, model emits "
    "### FILE: blocks, worker parses and writes them", style="List Number"
)
doc.add_paragraph(
    "7. Self-check pass: model compares written files vs agent rule, generates "
    "missing ones or reports LAYER_COMPLETE", style="List Number"
)
doc.add_paragraph(
    "8. Output artifacts zipped and available for download via UI", style="List Number"
)

doc.add_heading("3.3 One-Shot Mode (How RHDP Matches Cursor)", level=2)
doc.add_paragraph(
    "Instead of a tool loop (which costs $12-20/run due to per-call re-sending "
    "of the full conversation), RHDP pre-loads ALL context into a single prompt:"
)
t = doc.add_table(rows=5, cols=3)
t.style = "Light Grid Accent 1"
for i, h in enumerate(["Context", "Size", "How injected"]):
    t.rows[0].cells[i].text = h
rows = [
    ["Agent rule", "22,627 chars", "Full text in system prompt (no truncation)"],
    ["Context template", "45,747 chars", "Full text in system prompt (no truncation)"],
    ["Parsed PDFs", "Up to 60,000 chars", "In user prompt (Docling markdown)"],
    ["RDA/CIQ context", "Up to 20,000 chars", "In user prompt (prior stage output)"],
]
for r, row in enumerate(rows, 1):
    for c, val in enumerate(row):
        t.rows[r].cells[c].text = val

doc.add_paragraph(
    "\nThe model receives the same information as Cursor's agent (which reads it "
    "via tool calls), but in one shot — eliminating read round-trips."
)

doc.add_heading("3.4 Self-Check (Model Decides Completeness)", level=2)
doc.add_paragraph(
    "After Pass 1 (main generation), the worker sends a self-check prompt:\n\n"
    "  'SELF-CHECK — Review the Ansible layer against the Agent Rule. "
    "Compare every file the spec requires with what exists on disk. "
    "If ANY required files are missing, emit them. If the layer is COMPLETE, "
    "output exactly: LAYER_COMPLETE'\n\n"
    "This is equivalent to Cursor's tool loop where the agent can read what "
    "it wrote and decide whether to write more. No Python-side MIN_FILES or "
    "SMOKE hard-gates — the model decides."
)

doc.add_heading("3.5 Key Files", level=2)
t = doc.add_table(rows=6, cols=2)
t.style = "Light Grid Accent 1"
for i, h in enumerate(["File", "Purpose"]):
    t.rows[0].cells[i].text = h
files = [
    ["sprintweave-web/api/worker.py", "Core job execution — Docling, LLM calls, file parsing, self-check"],
    ["sprintweave-web/api/main.py", "FastAPI routes — POST /v1/jobs, GET status, resume, cancel"],
    ["sprintweave-web/ui/app.py", "Streamlit UI — file upload, pipeline picker, job table, log viewer"],
    ["sprintweave-web/llm/client.py", "LLM client — OpenAI-compatible API calls to MaaS/LiteLLM"],
    ["sprintweave-web/ui/pipelines.py", "Pipeline catalog — stages, agents, env-based filtering"],
]
for r, row in enumerate(files, 1):
    for c, val in enumerate(row):
        t.rows[r].cells[c].text = val

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════
# 4. Side-by-Side Comparison
# ═══════════════════════════════════════════════════════════════
doc.add_heading("4. Side-by-Side Comparison", level=1)

t = doc.add_table(rows=11, cols=3)
t.style = "Light Grid Accent 1"
for i, h in enumerate(["Aspect", "Cursor IDE", "RHDP / OpenShift"]):
    t.rows[0].cells[i].text = h
rows = [
    ["Execution", "Cursor SDK Agent.create()", "FastAPI worker thread"],
    ["LLM calls per layer", "~25 (read + write tool calls)", "2-3 (one-shot + self-check)"],
    ["How model gets spec", "read_file tool calls (3-5 calls)", "Pre-loaded in prompt (0 calls)"],
    ["How model writes files", "write_file tool calls (1 per file)", "### FILE: blocks in one response"],
    ["Who decides 'done'", "Model (stops calling tools)", "Model (says LAYER_COMPLETE)"],
    ["Spec injected", "Full (via read_file)", "Full (in system prompt, no truncation)"],
    ["Cost per run", "$20/month flat", "$4-6 per run (MaaS per-token)"],
    ["PDF processing", "Docling (same)", "Docling (same, shared converter)"],
    ["Output location", "Local filesystem (cpat-main/output/)", "PVC /data/jobs/<id>/workspace/output/"],
    ["Verification", "verify_paths — log only", "Self-check pass — model compares vs spec"],
]
for r, row in enumerate(rows, 1):
    for c, val in enumerate(row):
        t.rows[r].cells[c].text = val

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════
# 5. Docker Image Strategy
# ═══════════════════════════════════════════════════════════════
doc.add_heading("5. Docker Image Strategy (Two-Layer Build)", level=1)

doc.add_heading("5.1 Why Two Layers", level=2)
doc.add_paragraph(
    "Docling + HuggingFace models take 30-60 minutes to install and are ~2GB. "
    "Code changes happen daily. Splitting into two layers means code rebuilds "
    "take minutes, not hours."
)

doc.add_heading("5.2 Layer 1: Docling Base Image (rebuild rarely)", level=2)
doc.add_paragraph("File: sprintweave-web/Dockerfile.docling-base")
doc.add_paragraph(
    "  • Base: python:3.12-slim\n"
    "  • Installs: docling>=2.0.0, openpyxl, libgl1, libglib2.0\n"
    "  • Warms Docling import to fail early if install broken\n"
    "  • HF model cache at /opt/docling-cache (baked into image)\n"
    "  • Size: ~2.5GB\n"
    "  • Rebuild frequency: Only when Docling version changes"
)
doc.add_paragraph(
    "Build command:\n"
    "  podman build -f sprintweave-web/Dockerfile.docling-base \\\n"
    "    -t quay.io/rh_global_services_auto/sprintweave-docling-base:latest ."
)

doc.add_heading("5.3 Layer 2: App Image (rebuild often)", level=2)
doc.add_paragraph("File: sprintweave-web/Dockerfile.bundle")
doc.add_paragraph(
    "  • Base: Layer 1 (sprintweave-docling-base:latest)\n"
    "  • Installs: fastapi, uvicorn, streamlit, requests (thin deps)\n"
    "  • Copies: sprintweave-web/ (app code) + cpat-main/ (agent rules/templates)\n"
    "  • Build args: AGENTS_SRC (cpat-main or sprintweave-agents), PIPELINES\n"
    "  • Size: ~2.7GB (mostly base layer)\n"
    "  • Rebuild frequency: Every code or agent change (~3-5 minutes)"
)
doc.add_paragraph(
    "Build command:\n"
    "  podman build -f sprintweave-web/Dockerfile.bundle \\\n"
    "    --build-arg BASE_IMAGE=quay.io/.../sprintweave-docling-base:latest \\\n"
    "    --build-arg AGENTS_SRC=cpat-main \\\n"
    "    --build-arg PIPELINES=lab-requirements,rhoso \\\n"
    "    -t quay.io/rh_global_services_auto/sprintweave-rhoso:latest ."
)

doc.add_heading("5.4 Image Contents (inside the container)", level=2)
doc.add_paragraph(
    "/app/\n"
    "  sprintweave-web/         # App code (UI, API, LLM client)\n"
    "    api/worker.py           # Core job execution\n"
    "    api/main.py             # FastAPI routes\n"
    "    ui/app.py               # Streamlit UI\n"
    "    llm/client.py           # LLM client\n"
    "  sprintweave-agents/      # Agent rules + templates (from cpat-main)\n"
    "    .cursor/rules/          # agent-rhoso-da.md, agent-rhoso-deployment-engineer.md, ...\n"
    "    context/templates/      # rhoso-deployment-engineer.md, rhoso-da.md\n"
    "    context/pipelines.yaml  # Pipeline definitions\n"
    "/data/                     # PVC mount point\n"
    "  jobs/<job-id>/           # Per-job workspaces\n"
    "/opt/docling-cache/        # HuggingFace model cache"
)

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════
# 6. End-to-End Deployment Steps
# ═══════════════════════════════════════════════════════════════
doc.add_heading("6. End-to-End Deployment Steps", level=1)

doc.add_heading("6.1 Prerequisites", level=2)
doc.add_paragraph(
    "  • OpenShift cluster 4.14+ (RHDP lab or production)\n"
    "  • oc CLI logged in with project-create permissions\n"
    "  • podman login to quay.io\n"
    "  • Bastion host with SSH access (for RHDP)\n"
    "  • MaaS LLM API key (OpenAI-compatible endpoint)"
)

doc.add_heading("6.2 Step 1: Build Docling Base Image (one-time)", level=2)
doc.add_paragraph(
    "On bastion (or any machine with internet + podman):\n\n"
    "  podman build -f sprintweave-web/Dockerfile.docling-base \\\n"
    "    -t quay.io/rh_global_services_auto/sprintweave-docling-base:latest .\n"
    "  podman push quay.io/rh_global_services_auto/sprintweave-docling-base:latest\n\n"
    "This takes 30-60 minutes. Only rebuild when Docling version changes."
)

doc.add_heading("6.3 Step 2: Pack Sources (from developer laptop)", level=2)
doc.add_paragraph(
    "Script: sprintweave-web/scripts/pack-to-bastion.sh\n\n"
    "From WSL on the developer laptop:\n\n"
    "  cd ~/Documents/Projects/CFP_NEW\n"
    "  bash sprintweave-web/scripts/pack-to-bastion.sh\n\n"
    "This creates a tarball of sprintweave-web/ + cpat-main/ and SCPs it to the bastion.\n"
    "Environment variables: BASTION_HOST, BASTION_PORT."
)

doc.add_heading("6.4 Step 3: Build App Image + Push + Rollout (on bastion)", level=2)
doc.add_paragraph(
    "Script: sprintweave-web/scripts/bastion-build-rollout.sh\n\n"
    "On the bastion:\n\n"
    "  ssh -p 32158 lab-user@ssh.ocpv06.rhdp.net\n"
    "  mkdir -p ~/sprintweave-build\n"
    "  tar -xzf ~/sprintweave-bastion.tgz -C ~/sprintweave-build\n"
    "  bash ~/sprintweave-build/sprintweave-web/scripts/bastion-build-rollout.sh\n\n"
    "The script:\n"
    "  1. Detects agent tree layout (flat cpat-main or nested cpat-main/cpat-main)\n"
    "  2. podman build with Dockerfile.bundle (FROM docling-base, ~3-5 min)\n"
    "  3. podman push to quay.io\n"
    "  4. oc rollout restart deployment/sprintweave-api -n sprintweave\n\n"
    "Environment overrides: IMAGE, BASE_IMAGE, NAMESPACE, PIPELINES, SKIP_PUSH, SKIP_ROLLOUT."
)

doc.add_heading("6.5 Step 4: OpenShift Resources", level=2)
doc.add_paragraph(
    "The deployment needs:\n\n"
    "  • Namespace: sprintweave\n"
    "  • PVC: sprintweave-data (2Gi RWO) mounted at /data\n"
    "  • ConfigMap: LLM_PROVIDER, LLM_BASE_URL, LLM_MODEL, LLM_MAX_TOKENS\n"
    "  • Secret: LLM_API_KEY (MaaS bearer token)\n"
    "  • Deployment: sprintweave-api (1 replica, ports 8080+8000)\n"
    "  • Service + Route: expose port 8080 (Streamlit UI)\n\n"
    "YAML manifests are in: sprintweave-web/deploy/openshift/"
)

doc.add_heading("6.6 Step 5: Verify", level=2)
doc.add_paragraph(
    "  1. Open the Route URL in a browser → Streamlit UI loads\n"
    "  2. Upload a test PDF → job starts\n"
    "  3. Watch logs: Docling converts PDF, then LLM generates artifacts\n"
    "  4. Download output ZIP → compare against cpat-main/output/"
)

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════
# 7. Docling PDF Processing Pipeline
# ═══════════════════════════════════════════════════════════════
doc.add_heading("7. Docling PDF Processing Pipeline", level=1)

doc.add_paragraph(
    "Docling is an open-source document converter that extracts structured "
    "content from PDFs, DOCX, and PPTX files into clean Markdown."
)

doc.add_heading("7.1 Processing Flow", level=2)
doc.add_paragraph(
    "  1. User uploads PDF/DOCX/PPTX files via Streamlit UI\n"
    "  2. Files are saved to workspace/input/\n"
    "  3. _parse_inputs() creates a SHARED DocumentConverter (loaded once)\n"
    "  4. Each file is converted: PDF → Markdown\n"
    "  5. Output saved to workspace/parsed/<filename>.md\n"
    "  6. LLM stages read from parsed/*.md — never from raw PDFs"
)

doc.add_heading("7.2 Performance Optimization", level=2)
doc.add_paragraph(
    "The DocumentConverter loads HuggingFace layout models (~500MB) on first "
    "use. Previously, a new converter was created per file, causing the models "
    "to reload each time."
)
t = doc.add_table(rows=3, cols=3)
t.style = "Light Grid Accent 1"
for i, h in enumerate(["", "Before", "After"]):
    t.rows[0].cells[i].text = h
rows = [
    ["5 PDFs", "5 × model load (30s) + 5 × convert = ~200s", "1 × model load (30s) + 5 × convert = ~80s"],
    ["Speedup", "—", "~2.5× faster"],
]
for r, row in enumerate(rows, 1):
    for c, val in enumerate(row):
        t.rows[r].cells[c].text = val

doc.add_heading("7.3 Configuration", level=2)
t = doc.add_table(rows=5, cols=3)
t.style = "Light Grid Accent 1"
for i, h in enumerate(["Env Variable", "Default", "Purpose"]):
    t.rows[0].cells[i].text = h
rows = [
    ["SPRINTWEAVE_USE_DOCLING", "true", "Enable/disable Docling processing"],
    ["SPRINTWEAVE_DOCLING_OCR", "false", "Enable OCR for scanned PDFs (slower)"],
    ["SPRINTWEAVE_DOCLING_TABLES", "false", "Enable table structure recognition (needs extra HF models)"],
    ["DOCLING_CACHE_DIR", "/opt/docling-cache", "HuggingFace model cache directory"],
]
for r, row in enumerate(rows, 1):
    for c, val in enumerate(row):
        t.rows[r].cells[c].text = val

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════
# 8. LLM Integration
# ═══════════════════════════════════════════════════════════════
doc.add_heading("8. LLM Integration (MaaS / LiteLLM)", level=1)

doc.add_paragraph(
    "SprintWeave calls an OpenAI-compatible /v1/chat/completions endpoint. "
    "On RHDP, this is MaaS (Model as a Service) with LiteLLM proxy."
)

doc.add_heading("8.1 Configuration", level=2)
t = doc.add_table(rows=6, cols=3)
t.style = "Light Grid Accent 1"
for i, h in enumerate(["Env Variable", "Example", "Purpose"]):
    t.rows[0].cells[i].text = h
rows = [
    ["LLM_PROVIDER", "openai", "Client type (openai-compatible)"],
    ["LLM_BASE_URL", "https://maas-rhdp.apps.../v1", "API endpoint"],
    ["LLM_MODEL", "claude-opus-4-6", "Model name"],
    ["LLM_API_KEY", "sk-Pga...", "Bearer token (NOT sk-ant- format)"],
    ["LLM_MAX_TOKENS", "16384", "Max output tokens per call"],
]
for r, row in enumerate(rows, 1):
    for c, val in enumerate(row):
        t.rows[r].cells[c].text = val

doc.add_heading("8.2 Call Pattern (One-Shot Mode)", level=2)
doc.add_paragraph(
    "For each RDE sub-stage (Ansible, Backend, Frontend, Docker):\n\n"
    "  System prompt = Instructions + Agent Rule (22k) + Context Template (45k)\n"
    "  User prompt = Sub-stage prompt + RDA context (20k) + Parsed PDFs (60k)\n"
    "  Total input: ~150k chars ≈ ~40k tokens\n\n"
    "  Pass 1: Model generates all files as ### FILE: blocks\n"
    "  Pass 2: Self-check — model compares output vs spec\n"
    "  Pass 3: (if needed) Model generates missing files\n"
    "  Pass N: Model says LAYER_COMPLETE → stop\n\n"
    "Up to RDE_LAYER_MAX_CALLS (default 4) passes per layer."
)

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════
# 9. Agent Rule & Context Template System
# ═══════════════════════════════════════════════════════════════
doc.add_heading("9. Agent Rule & Context Template System", level=1)

doc.add_paragraph(
    "The agent rule and context template together form the complete spec that "
    "drives code generation. Both Cursor and RHDP use the same files."
)

doc.add_heading("9.1 Agent Rule (.cursor/rules/agent-rhoso-deployment-engineer.md)", level=2)
doc.add_paragraph(
    "306 lines, 22,627 chars. Contains:\n\n"
    "  • When This Activates — trigger conditions\n"
    "  • Primary Deliverables — table of output paths\n"
    "  • Multi-Pass Generation — sub-stage table (Ansible/Backend/Frontend/Docker)\n"
    "  • Manifest Coverage — 19 Jinja templates with CR kinds and conditionals\n"
    "  • Deployment Task Files — 25 tasks in 9 subdirectories with purpose and conditionals\n"
    "  • Key Extensions — v1.4 additions (Trident, LSO, Observability, Multipath)\n"
    "  • Frontend Components — 11 named components (RhosoInstallation, Layout, Login, etc.)\n"
    "  • Backend Routes — 18 API routes with methods and purposes"
)

doc.add_heading("9.2 Context Template (context/templates/rhoso-deployment-engineer.md)", level=2)
doc.add_paragraph(
    "1014 lines, 45,747 chars. Contains:\n\n"
    "  • RHOSO Deployment Documentation Mapping — chapters to CIQ tables\n"
    "  • CPAT Source Code Generation — Pattern G (project blueprint with 9-subdir layout)\n"
    "  • Pattern H — Ansible task patterns (include_tasks, template rendering, oc wait)\n"
    "  • Pattern I — Jinja template-to-CIQ variable mapping (19 templates)\n"
    "  • Pattern J — Flask backend patterns (routes, streaming, error handling)\n"
    "  • Pattern K — React frontend patterns (component structure, API calls)\n"
    "  • Pattern L — Docker/bringup patterns (multi-stage build, Podman lifecycle)"
)

doc.add_heading("9.3 How They Work Together", level=2)
doc.add_paragraph(
    "The agent rule says WHAT to generate (file list, component names, template names). "
    "The context template says HOW to generate it (code patterns, variable mappings, "
    "project structure). Both are injected into every LLM call — in Cursor via "
    "read_file, in RHDP via pre-loading in the system prompt."
)

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════
# 10. Cost Analysis
# ═══════════════════════════════════════════════════════════════
doc.add_heading("10. Cost Analysis", level=1)

t = doc.add_table(rows=4, cols=5)
t.style = "Light Grid Accent 1"
for i, h in enumerate(["Approach", "LLM Calls/Run", "Cost/Run", "Match Rate", "Notes"]):
    t.rows[0].cells[i].text = h
rows = [
    ["Cursor tool-loop", "~100 (25/layer × 4)", "$20/mo flat", "~95%", "Unlimited within subscription"],
    ["RHDP one-shot (old)", "4-8", "$2-4", "~64%", "Template truncated to 8k chars"],
    ["RHDP one-shot (current)", "8-12", "$4-6", "~90%+", "Full spec + self-check + 9-subdir prompt"],
]
for r, row in enumerate(rows, 1):
    for c, val in enumerate(row):
        t.rows[r].cells[c].text = val

doc.add_paragraph(
    "\nThe cost difference comes from the number of LLM round-trips. Each tool-loop "
    "step re-sends the entire conversation history. One-shot sends the full context "
    "once, and the model outputs all files in a single response."
)

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════
# 11. File Tree
# ═══════════════════════════════════════════════════════════════
doc.add_heading("11. File Tree — Cursor vs SprintWeave Output", level=1)

doc.add_heading("11.1 CPAT Output Structure (Cursor / cpat-main)", level=2)
doc.add_paragraph(
    "output/cpat/\n"
    "  ansible/\n"
    "    ansible.cfg\n"
    "    playbooks/\n"
    "      main.yml\n"
    "      roles/\n"
    "        register/tasks/rhoso.yml\n"
    "        rhoso/\n"
    "          defaults/main.yml\n"
    "          tasks/\n"
    "            main.yml\n"
    "            common/ (_load_ciq_context, _validate_ciq, _verify_required_crds)\n"
    "            prerequisites/ (prerequisites, prerequisite_operator_installation, disconnected_environment_preparation)\n"
    "            secrets/ (generate_osp_secret, generate_data_plane_secrets, generate_ceph_secret)\n"
    "            storage/ (persistent_storage_configuration, trident_storage_configuration, _generate_storage_crs)\n"
    "            network/ (network_preparation)\n"
    "            control_plane/ (5 files)\n"
    "            data_plane/ (2 files)\n"
    "            post_deployment/ (2 files)\n"
    "            manifests/ (2 files)\n"
    "          templates/ (19 Jinja2 .yaml.j2 files)\n"
    "    library/ (parse_ciq.py, get_ciq_validation_schema.py)\n"
    "    module_utils/ (cpat_ciq_parser.py, rhoso_ciq_schema.py)\n"
    "  backend/ (app.py, ansible_runner.py, ciq_parser.py, models.py, paths.py, ...)\n"
    "  frontend/\n"
    "    src/components/ (11 JSX: RhosoInstallation, Layout, Login, Dashboard, ...)\n"
    "    src/ (App.jsx, api.js, index.css, main.jsx)\n"
    "    package.json, vite.config.js, index.html\n"
    "    nginx/default.conf\n"
    "  docker/entrypoint.sh\n"
    "  Dockerfile, bringup.sh, USAGE.md, .dockerignore"
)

doc.add_heading("11.2 File Counts", level=2)
t = doc.add_table(rows=5, cols=3)
t.style = "Light Grid Accent 1"
for i, h in enumerate(["Layer", "Cursor (cpat-main)", "RHDP target"]):
    t.rows[0].cells[i].text = h
rows = [
    ["Ansible (RHOSO)", "~50 files", "~50 files (9-subdir layout)"],
    ["Backend", "11 files", "10-11 files"],
    ["Frontend", "~22 files (RHOSO)", "~22 files"],
    ["Docker + root", "~8 files", "~8 files"],
]
for r, row in enumerate(rows, 1):
    for c, val in enumerate(row):
        t.rows[r].cells[c].text = val

# ── Save ───────────────────────────────────────────────────────
doc.save(str(OUT))
print(f"Saved: {OUT}")
