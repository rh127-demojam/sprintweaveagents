"""Convert PRESENTATION-E2E-ARCHITECTURE.md → .docx"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor

HERE = Path(__file__).resolve().parent
MD_PATH = HERE / "PRESENTATION-E2E-ARCHITECTURE.md"
OUT_PATH = HERE / "PRESENTATION-E2E-ARCHITECTURE.docx"

INLINE_RE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")


def add_runs(paragraph, content: str) -> None:
    for part in INLINE_RE.split(content):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
        else:
            paragraph.add_run(part)


def main() -> None:
    text = MD_PATH.read_text(encoding="utf-8")
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    lines = text.splitlines()
    i = 0
    in_code = False
    code_buf: list[str] = []
    table_buf: list[str] = []

    def flush_table() -> None:
        nonlocal table_buf
        if not table_buf:
            return
        rows: list[list[str]] = []
        for row in table_buf:
            if re.match(r"^\|?\s*-+", row):
                continue
            cells = [c.strip() for c in row.strip().strip("|").split("|")]
            rows.append(cells)
        table_buf = []
        if not rows:
            return
        cols = max(len(r) for r in rows)
        t = doc.add_table(rows=len(rows), cols=cols)
        t.style = "Table Grid"
        for ri, row in enumerate(rows):
            for ci in range(cols):
                cell = t.rows[ri].cells[ci]
                cell.text = row[ci] if ci < len(row) else ""
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(9)
                        run.font.name = "Calibri"
                        if ri == 0:
                            run.bold = True
        doc.add_paragraph()

    def flush_code() -> None:
        nonlocal code_buf
        if not code_buf:
            return
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(8)
        run = p.add_run("\n".join(code_buf))
        run.font.name = "Consolas"
        run.font.size = Pt(8)
        code_buf = []

    while i < len(lines):
        line = lines[i]
        if line.startswith("```"):
            if in_code:
                flush_code()
                in_code = False
            else:
                flush_table()
                in_code = True
                code_buf = []
            i += 1
            continue
        if in_code:
            code_buf.append(line)
            i += 1
            continue
        if line.strip().startswith("|"):
            table_buf.append(line)
            i += 1
            continue
        flush_table()

        if not line.strip():
            i += 1
            continue
        if line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=0)
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=2)
        elif line.startswith("#### "):
            doc.add_heading(line[5:].strip(), level=3)
        elif line.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            add_runs(p, line[2:].strip())
            for run in p.runs:
                run.italic = True
        elif line.startswith("- ") or line.startswith("* "):
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, line[2:].strip())
        elif re.match(r"^\d+\.\s", line):
            p = doc.add_paragraph(style="List Number")
            add_runs(p, re.sub(r"^\d+\.\s", "", line))
        elif line.strip() == "---":
            pass
        else:
            p = doc.add_paragraph()
            add_runs(p, line.strip())
        i += 1

    flush_table()
    flush_code()
    doc.save(OUT_PATH)
    print(OUT_PATH)
    print("bytes", OUT_PATH.stat().st_size)


if __name__ == "__main__":
    main()
